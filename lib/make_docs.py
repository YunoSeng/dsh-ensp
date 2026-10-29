"""Render the two lab deliverables with python-docx.

Usage: python make_docs.py <manifest.json>

Manifest shape::

    {
      "title": "实训任务3-1 知识拓展",
      "subtitle": "姓名 / 班级 / 日期",
      "mode": "captioned" | "bare" | "template",
      "org": "华为 eNSP 实验报告",
      "sections": [
        {"heading": "一、配置 SW1", "paragraphs": ["..."],
         "images": [{"path": "...png", "caption": "图1 配置 SW1"}]}
      ],
      "credentials": [{"item": "console 登录口令", "value": "Huawei@123"}],
      "notes": ["..."]
    }

``captioned`` writes a short explanation under every screenshot; ``bare`` keeps
only the task heading and the screenshot itself, with no prose at all.

``template`` (bare copy only) takes the teacher's original .docx, keeps every
character of it, and drops each section's screenshots into the blank paragraphs
that follow its ``步骤N`` heading — the shape a student hands in.
"""

import copy
import json
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

BODY_FONT = "宋体"
CODE_FONT = "Consolas"
MAX_IMAGE_WIDTH_CM = 15.5
STEP_RE = re.compile(r"步骤\s*([0-9]+)")


def set_style_font(style, name, size=None, bold=None):
    style.font.name = name
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def add_body(doc, text, *, bold=False, code=False, size=10.5, italic=False):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    run.font.name = CODE_FONT if code else BODY_FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    return paragraph


def add_image(doc, path, caption, *, captioned, width_cm=MAX_IMAGE_WIDTH_CM):
    if not os.path.exists(path):
        raise SystemExit("missing image: %s" % path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run().add_picture(path, width=Cm(width_cm))
    if captioned and caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)
        run = cap.add_run(caption)
        run.font.name = BODY_FONT
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)


def _is_blank(paragraph):
    """True for a paragraph that holds neither text nor a picture."""
    if paragraph.text.strip():
        return False
    if paragraph._p.findall('.//' + qn('w:drawing')):
        return False
    if paragraph._p.findall('.//' + qn('w:pict')):
        return False
    return True


def _anchor_key(section):
    """Map a section to the line in the template it belongs under."""
    explicit = section.get("anchor")
    if explicit:
        return explicit
    heading = section.get("heading", "")
    match = STEP_RE.search(heading)
    if match:
        return "步骤" + match.group(1)
    return heading.strip()


def fill_template(template_path, out_path, sections, width_cm=13.5):
    """Copy the teacher's document and drop screenshots into its blank slots."""
    if not os.path.exists(template_path):
        raise SystemExit("missing template: %s" % template_path)
    doc = Document(template_path)
    paragraphs = doc.paragraphs

    wanted = []
    for section in sections:
        key = _anchor_key(section)
        images = [image for image in section.get("images", []) if image.get("path")]
        if key and images:
            wanted.append((key, images, section.get("width_cm", width_cm)))

    index = {}
    for i, paragraph in enumerate(paragraphs):
        lines = [line.strip() for line in paragraph.text.split("\n")]
        for key, _images, _width in wanted:
            if key in index:
                continue
            if any(line.startswith(key) for line in lines):
                index[key] = i
    missing = [key for key, _i, _w in wanted if key not in index]
    if missing:
        raise SystemExit("template has no paragraph starting with: %s" % ", ".join(missing))

    order = sorted(((key, index[key], images, width) for key, images, width in wanted),
                   key=lambda item: item[1])
    for n, (key, start, images, width) in enumerate(order):
        end = order[n + 1][1] if n + 1 < len(order) else len(paragraphs)
        blanks = [p for p in paragraphs[start + 1:end] if _is_blank(p)]
        anchor = paragraphs[start]
        for j, image in enumerate(images):
            path = image["path"]
            if not os.path.exists(path):
                raise SystemExit("missing image: %s" % path)
            if j < len(blanks):
                target = blanks[j]
            else:
                # Ran out of blank lines: clone the formatting of the last one.
                new_p = OxmlElement('w:p')
                source = blanks[-1]._p if blanks else None
                if source is not None:
                    ppr = source.find(qn('w:pPr'))
                    if ppr is not None:
                        new_p.append(copy.deepcopy(ppr))
                anchor._p.addnext(new_p)
                target = Paragraph(new_p, anchor._parent)
            anchor = target
            for run in list(target.runs):
                run._r.getparent().remove(run._r)
            target.alignment = WD_ALIGN_PARAGRAPH.CENTER
            target.add_run().add_picture(path, width=Cm(width))

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    if os.path.exists(out_path):
        os.remove(out_path)
    doc.save(out_path)
    return {"output": out_path, "anchors": len(order),
            "images": sum(len(images) for _k, _i, images, _w in order)}


def build(manifest, out_path):
    if manifest.get("mode") == "template":
        if not manifest.get("template"):
            raise SystemExit("template mode needs a template path")
        return fill_template(manifest["template"], out_path, manifest.get("sections", []),
                             manifest.get("width_cm", 13.5))
    captioned = manifest.get("mode", "captioned") == "captioned"
    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)

    set_style_font(doc.styles["Normal"], BODY_FONT, 10.5)
    for name in ("Heading 1", "Heading 2", "Title"):
        if name in doc.styles:
            set_style_font(doc.styles[name], BODY_FONT, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(manifest.get("title", "实验报告"))
    run.font.name = BODY_FONT
    run.font.size = Pt(18)
    run.bold = True
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    if captioned and manifest.get("subtitle"):
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = sub.add_run(manifest["subtitle"])
        run.font.name = BODY_FONT
        run.font.size = Pt(10.5)
        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    for item in manifest.get("sections", []):
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(14)
        heading.paragraph_format.space_after = Pt(6)
        run = heading.add_run(item.get("heading", ""))
        run.font.name = BODY_FONT
        run.font.size = Pt(13)
        run.bold = True
        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

        if captioned:
            for text in item.get("paragraphs", []):
                add_body(doc, text)
            for command in item.get("commands", []):
                add_body(doc, command, code=True, size=9.5)

        for image in item.get("images", []):
            add_image(doc, image["path"], image.get("caption", ""),
                      captioned=captioned,
                      width_cm=image.get("width_cm", MAX_IMAGE_WIDTH_CM))

    if captioned and manifest.get("credentials"):
        heading = doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(16)
        run = heading.add_run("附：本次实验账号与口令对照表")
        run.font.name = BODY_FONT
        run.font.size = Pt(13)
        run.bold = True
        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text = "项目"
        header[1].text = "当前生效值"
        for row in manifest["credentials"]:
            cells = table.add_row().cells
            cells[0].text = str(row.get("item", ""))
            cells[1].text = str(row.get("value", ""))
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
                        run.font.name = BODY_FONT
                        run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    if captioned:
        for note in manifest.get("notes", []):
            add_body(doc, note, size=9.5, italic=True)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    if os.path.exists(out_path):
        os.remove(out_path)
    doc.save(out_path)
    return out_path


def main():
    if len(sys.argv) < 2:
        raise SystemExit("usage: make_docs.py <manifest.json>")
    with open(sys.argv[1], "r", encoding="utf-8") as handle:
        manifest = json.load(handle)
    out = manifest.get("output")
    if not out:
        raise SystemExit("manifest needs an output path")
    built = build(manifest, out)
    if isinstance(built, dict):
        print(json.dumps(built, ensure_ascii=False))
        return
    print(json.dumps({"output": out, "sections": len(manifest.get("sections", [])),
                      "images": sum(len(s.get("images", [])) for s in manifest.get("sections", []))},
                     ensure_ascii=False))


if __name__ == "__main__":
    main()
